"""Write the Word review sheet for one model's part details: what a learner would read for each kind of part, the evidence behind each
line, and the catalogue cross-check, each beside an empty comments column.

The sheet is generated from the part store (knowledge/part-cards/parts, docs/part-store.md) and is not committed; run this again after the
store changes. Every definition stays 'draft' until you have read it. Part numbers are left off on purpose. The IO-520 catalogue rows are kept
out of Git (the repository is public); when the private copy is present (data/gtsio_part_authoring/private/definitions-engineering.json,
written by scripts/gtsio_part_authoring/private_check.py) the cross-check column also names the catalogue page, figure and item so you can find
each row. Without it the sheet shows the alignment result, its note and the GTSIO manual figure keys.

    python scripts/build_part_review_sheet.py [--model gtsio520-h-r10-mechanism-study | cylinder] [--out "../Lesson Plans/Review sheets"]

For the operating cylinder the sheet ends with what became of every value in the retired hand-written table.
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from build_review_sheets import GREY, HEADER_FILL, TEAL, set_widths, shade

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from cad_pipeline.intent import part_store  # noqa: E402

PRIVATE_DEFINITIONS = ROOT / "data/gtsio_part_authoring/private/definitions-engineering.json"
R10 = "gtsio520-h-r10-mechanism-study"

SHEETS = {
    R10: dict(
        file="GTSIO-520-H part details - draft review.docx",
        heading="GTSIO-520-H part details: draft text for your review",
        card="the GTSIO-520-H whole-engine card",
    ),
    "cylinder": dict(
        file="Operating cylinder part details - draft review.docx",
        heading="Operating cylinder part details: draft text for your review",
        card="the operating cylinder card",
    ),
}

STATUS_WORDS = {
    "specified": "stated in a source",
    "derived": "worked out from a source",
    "inferred": "reasoned, not in a source",
    "illustrative": "learning prompt",
    "unsourced": "no source found",
}
ALIGNMENT_WORDS = {
    "confirmed": "confirmed: the IO-520 catalogue and the GTSIO manual agree",
    "gtsio_only": "GTSIO only: in the GTSIO manual, not in the IO-520 catalogue",
    "io520_only": "IO-520 only: in the catalogue, not found in the GTSIO manual",
    "mismatch": "MISMATCH: the two books disagree",
    "not_found": "not found in either book",
}
RED = RGBColor(0xB4, 0x23, 0x18)


def run(paragraph, text: str, *, bold=False, italic=False, size=9, color: RGBColor | None = None):
    piece = paragraph.add_run(text)
    piece.bold = bold
    piece.italic = italic
    piece.font.size = Pt(size)
    if color is not None:
        piece.font.color.rgb = color
    return piece


def line(cell, label: str | None, text: str, *, size=9, color: RGBColor | None = None, first=False):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    if label:
        run(paragraph, f"{label}: ", bold=True, size=size, color=color)
    run(paragraph, text, size=size, color=color)
    return paragraph


def evidence_text(claim: dict) -> str:
    where = claim.get("source") or claim.get("basis") or ""
    return f"{STATUS_WORDS[claim['status']]}. {where}"


def ipc_text(row: dict) -> str:
    flag = "" if row.get("ocrClean", True) else " (OCR doubtful: check the page image)"
    return f"page {row['page']}, figure {row['figure']}, item {row['index']}: {row['description']}{flag}"


def manual_text(row: dict) -> str:
    return f"{row['figure']} callout {row['callout']}: {row['key']}"


def title_of(definition: dict) -> str:
    student = definition["student"]
    return student["title"]["text"] + (f" ({student['context']})" if student.get("context") else "")


def part_cell(cell, definition: dict, members: list[tuple[str, str]], total: int) -> None:
    student = definition["student"]
    line(cell, None, title_of(definition), size=11, color=TEAL, first=True).runs[0].bold = True
    line(cell, None, f"{total} in the model · {definition['id']}", size=8, color=GREY)
    if student["aliases"]:
        line(cell, "Also called", ", ".join(student["aliases"]), size=8, color=GREY)
    old = definition["engineering"].get("priorData", {}).get("labels", [])
    if old and total > 6:
        line(cell, "Was", old[0], size=8, color=GREY)
    if members and total <= 6:
        line(cell, "In the Parts list", "; ".join(name for _, name in members), size=8, color=GREY)


def learner_cell(cell, definition: dict, titles: dict[str, str], own: dict[str, dict]) -> None:
    student = definition["student"]
    line(cell, "Does", student["function"]["text"], first=True)
    line(cell, "Made of", student["madeOf"]["text"])
    line(cell, "Sits", student["sits"]["text"])
    line(cell, "Works with", ", ".join(titles[i] for i in student["worksWith"]["ids"]))
    line(cell, "Look for", student["lookFor"]["text"])
    if student.get("catalogue"):
        line(cell, "Catalogue line", f"{student['catalogue']['description']}. {student['catalogue']['note']}", size=8, color=GREY)
    for name, claims in own.items():
        for field, claim in claims.items():
            line(cell, f"{name}, in this model", f"{field}: {claim['text']}", size=8, color=TEAL)


def evidence_cell(cell, definition: dict, own: dict[str, dict]) -> None:
    student = definition["student"]
    engineering = definition["engineering"]
    line(cell, "Name", evidence_text(student["title"]), size=8, first=True)
    for key, label in [("function", "Does"), ("madeOf", "Made of"), ("sits", "Sits"), ("worksWith", "Works with")]:
        line(cell, label, evidence_text(student[key]), size=8)
    for name, claims in own.items():
        for field, claim in claims.items():
            line(cell, f"{name} {field}", evidence_text(claim), size=8, color=TEAL)
    alignment = engineering["alignment"]
    mark = RED if alignment["result"] in ("mismatch", "not_found") else GREY
    line(cell, "Cross-check", f"{ALIGNMENT_WORDS[alignment['result']]}. {alignment['note']}", size=8, color=mark)
    for row in engineering.get("ipcIo520", []):
        line(cell, "IO-520 catalogue", ipc_text(row), size=8, color=GREY)
    for row in engineering["gtsioManual"]:
        line(cell, "GTSIO manual", manual_text(row), size=8, color=GREY)


def parts_list_names(instance_doc: dict) -> dict[str, str]:
    """Part id -> the name the app's Parts list gives it (the cylinder's catalogue; other models use the id)."""
    if instance_doc["idsFrom"]["file"] != "web/components.json":
        return {}
    catalogue = json.loads((ROOT / "web/components.json").read_text(encoding="utf-8"))
    return {p["cad_stable_id"]: p["display_name"] for p in catalogue["parts"]}


def legacy_table(document: Document, definitions: list[dict], names: dict[str, str]) -> None:
    items = [(item, d) for d in definitions for item in d["engineering"].get("legacyValues", [])]
    if not items:
        return
    document.add_page_break()
    document.add_heading(f"What became of the old hand-written table ({len(items)} values)", level=2)
    intro = document.add_paragraph()
    run(intro, "Before this change the cylinder's part card read src/data/componentMetadata.ts: for 28 parts, a material line and an operating line, with alloy specifications "
               "and temperatures and no source. Each value was checked against the GTSIO-520 overhaul manual, the IO-520 maintenance manual and the lesson text. "
               "'Kept' means a source states it and the new card says it in plain words. 'No source found' means no source read supports it, so it is not on the card; "
               "where part of a value is supported or a source contradicts it, the note says so. If you know a source for any of these, say so in the last column and it can come back.", size=10)
    counts = Counter(item["status"] for item, _ in items)
    summary = document.add_paragraph()
    run(summary, "Result: " + ", ".join(f"{counts[s]} {STATUS_WORDS[s]}" for s in ("specified", "derived", "unsourced") if counts[s]) + ".", bold=True, size=10)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, head in zip(table.rows[0].cells, ["Part (Parts list)", "Old value", "What became of it", "Source, or what was found", "Your comments"]):
        shade(cell, HEADER_FILL)
        run(cell.paragraphs[0], head, bold=True, size=9)
    for item, _ in sorted(items, key=lambda pair: (names.get(pair[0]["part"], pair[0]["part"]).lower(), pair[0]["field"])):
        cells = table.add_row().cells
        line(cells[0], None, names.get(item["part"], item["part"]), size=9, first=True)
        line(cells[0], None, "material" if item["field"] == "material" else "operating conditions", size=8, color=GREY)
        line(cells[1], None, item["text"], size=9, first=True)
        verdict = "kept" if item["disposition"] == "kept" else "not on the card"
        line(cells[2], None, f"{STATUS_WORDS[item['status']]}; {verdict}", size=9, first=True, color=RED if item["status"] == "unsourced" else None)
        line(cells[3], None, item.get("source") or item.get("basis") or "", size=8, first=True)
    set_widths(table, [4.0, 6.2, 3.6, 9.6, 3.3])


def build(model: str, out_dir: Path) -> Path:
    spec = SHEETS[model]
    loaded = part_store.load_store()
    instance_doc = next(doc for doc in loaded["instances"] if doc["model"] == model)
    part_ids = part_store.model_part_ids(instance_doc)
    hits, unmatched, ambiguous = part_store.resolve_instances(instance_doc, part_ids)
    if unmatched or ambiguous:
        raise SystemExit(f"{model}: parts without exactly one definition; run: python -m cad_pipeline.intent.part_store")
    explicit = {row["id"]: row for row in instance_doc.get("parts", [])}
    names = parts_list_names(instance_doc)

    source = loaded["definitions"]
    if PRIVATE_DEFINITIONS.exists():
        private = {d["id"]: d for d in json.loads(PRIVATE_DEFINITIONS.read_text(encoding="utf-8"))["definitions"]}
        source = [private.get(d["id"], d) for d in loaded["definitions"]]
    by_id = {d["id"]: d for d in source}
    titles = {d["id"]: title_of(d) for d in source}
    used = [by_id[i] for i in dict.fromkeys(hits.values())]
    members: dict[str, list[tuple[str, str]]] = {}
    for pid, did in hits.items():
        members.setdefault(did, []).append((pid, names.get(pid, pid)))

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    document.add_heading(spec["heading"], level=1)
    intro = document.add_paragraph()
    run(intro, f"Tapping a part in {spec['card']} shows a plain title, one sentence on what the part does, what it is made of, where it sits, "
               "what it works with, something to look for, and the catalogue wording. Every line below is a draft and stays marked 'Draft text' for students until you review it. "
               "Write in the last column; nothing here has been shown to students as reviewed.", size=10)
    rules = document.add_paragraph()
    run(rules, "How the text was written. ", bold=True, size=10)
    run(rules, "The title is the part noun from the IO-520 parts catalogue, in plain order and four words at most; where the catalogue has no row the GTSIO manual's figure key gives the name. "
               "Sizes and specifications are not in the plain lines: the 'catalogue line' carries the catalogue wording and is tagged as IO-520 (a similar engine) unless the GTSIO manual confirms it. "
               "No part numbers are shown to students (the GTSIO-520-H catalogue, Form X-30046A, is not available). Numbers appear only where a lesson or the manual gives them. "
               "Each line carries its evidence: 'stated in a source' and 'worked out from a source' name the source; 'reasoned, not in a source' gives the basis and is a claim to check; "
               "'look for' lines are prompts for the learner and claim nothing about the engine.", size=10)
    if model == "cylinder":
        shared = Counter(1 for did in set(hits.values()) if any(m for m in loaded["instances"] if m["model"] == R10 and did in {r["definition"] for r in m.get("rules", [])}))
        reused = sum(shared.values())
        also = document.add_paragraph()
        run(also, "Shared with the whole-engine card. ", bold=True, size=10)
        run(also, f"{reused} of the {len(used)} kinds of part below are the same definitions the GTSIO-520-H whole-engine card reads; the cylinder adds {len(used) - reused} of its own "
                  "(the crankshaft throw, the rocker housing and the fuel discharge nozzle) and gives its own wording for a few parts where the shared line would name the whole engine's model. "
                  "Those lines are shown in teal. Changing a shared definition changes both cards.", size=10)

    statuses = Counter()
    for definition in used:
        for key in ("function", "madeOf", "sits"):
            statuses[(key, definition["student"][key]["status"])] += 1
    alignment = Counter(d["engineering"]["alignment"]["result"] for d in used)
    summary = document.add_paragraph()
    run(summary, f"{len(used)} kinds of part cover all {len(part_ids):,} modelled parts. ", bold=True, size=10)
    run(summary, "Function lines: " + ", ".join(f"{statuses[('function', s)]} {STATUS_WORDS[s]}" for s in ("specified", "derived", "inferred")) + ". "
               "Catalogue cross-check: " + ", ".join(f"{alignment[a]} {a.replace('_', ' ')}" for a in ("confirmed", "gtsio_only", "io520_only", "mismatch", "not_found") if alignment[a]) + ".", size=10)

    flagged = [d for d in used if d["engineering"]["alignment"]["result"] in ("mismatch", "not_found", "io520_only")]
    if flagged:
        document.add_heading("Where the two books disagree or one is silent", level=2)
        note = document.add_paragraph()
        run(note, "These are the entries to look at first. Each is in the table below with the full note.", size=10)
        for definition in flagged:
            paragraph = document.add_paragraph(style="List Bullet")
            run(paragraph, f"{titles[definition['id']]} ", bold=True, size=9)
            run(paragraph, f"[{definition['engineering']['alignment']['result'].replace('_', ' ')}] {definition['engineering']['alignment']['note']}", size=9)

    by_system: dict[str, list[dict]] = {}
    for definition in used:
        by_system.setdefault(definition["student"]["system"], []).append(definition)
    for system in loaded["store"]["systems"]:
        group = by_system.get(system["id"], [])
        if not group:
            continue
        document.add_page_break()
        document.add_heading(f"{system['label']} ({len(group)})", level=2)
        question = document.add_paragraph()
        run(question, system["question"], italic=True, size=10, color=GREY)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, head in zip(table.rows[0].cells, ["Part", "What a learner would read", "Where it comes from", "Your comments"]):
            shade(cell, HEADER_FILL)
            run(cell.paragraphs[0], head, bold=True, size=9)
        for definition in sorted(group, key=lambda d: titles[d["id"]].lower()):
            cells = table.add_row().cells
            mine = members[definition["id"]]
            # Wording this model gives a part itself, by the part's name in the Parts list.
            own = {names.get(pid, pid): explicit[pid]["claims"] for pid, _ in mine if pid in explicit and explicit[pid].get("claims")}
            part_cell(cells[0], definition, mine, len(mine))
            learner_cell(cells[1], definition, titles, own)
            evidence_cell(cells[2], definition, own)
        set_widths(table, [3.6, 9.2, 9.6, 4.3])

    if model == "cylinder":
        legacy_table(document, used, names)

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / spec["file"]
    document.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--model", default=R10, choices=sorted(SHEETS), help="the model whose part text to write out")
    parser.add_argument("--out", default=str(ROOT.parent / "Lesson Plans" / "Review sheets"), help="folder for the sheet")
    args = parser.parse_args()
    print(build(args.model, Path(args.out)))


if __name__ == "__main__":
    main()
