"""Write the Word review sheet for the GTSIO-520-H part details: what a learner would read for each kind of part, the evidence
behind each line, and the catalogue cross-check, each beside an empty comments column.

The sheet is generated from data/gtsio520-part-definitions.json and is not committed; run this again after the definitions change.
Every definition stays 'draft' until you have read it. Part numbers are left off on purpose: the cross-check names the catalogue page,
figure and index so you can find the row, and the numbers stay in the data file for engineering use.

    python scripts/build_gtsio_part_review_sheet.py [--out "../Lesson Plans/Review sheets"]
"""
from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from build_review_sheets import GREY, HEADER_FILL, TEAL, set_widths, shade

ROOT = Path(__file__).resolve().parent.parent
DEFINITIONS = ROOT / "data/gtsio520-part-definitions.json"
FILE_NAME = "GTSIO-520-H part details - draft review.docx"

STATUS_WORDS = {
    "specified": "stated in a source",
    "derived": "worked out from a source",
    "inferred": "reasoned, not in a source",
    "illustrative": "learning prompt",
}
ALIGNMENT_WORDS = {
    "confirmed": "confirmed: the IO-520 catalogue and the GTSIO manual agree",
    "gtsio_only": "GTSIO only: in the GTSIO manual, not in the IO-520 catalogue",
    "io520_only": "IO-520 only: in the catalogue, not found in the GTSIO manual",
    "mismatch": "MISMATCH: the two books disagree",
    "not_found": "not found in either book",
}


def run(paragraph, text: str, *, bold=False, italic=False, size=9, color: RGBColor | None = None):
    piece = paragraph.add_run(text)
    piece.bold = bold
    piece.italic = italic
    piece.font.size = Pt(size)
    if color is not None:
        piece.font.color.rgb = color
    return piece


def line(cell, label: str | None, text: str, *, size=9, color: RGBColor | None = None, first=False):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2)
    if label:
        run(paragraph, f"{label}: ", bold=True, size=size, color=color)
    run(paragraph, text, size=size, color=color)
    return paragraph


def evidence_text(claim: dict) -> str:
    where = claim.get("source") or claim.get("basis") or ""
    return f"{STATUS_WORDS[claim['status']]}. {where}"


def ipc_text(row: dict) -> str:
    flag = "" if row.get("ocrClean", True) else " (OCR doubtful: check the page image)"
    return f"page {row['page']}, figure {row['figure']}, item {row['index']}: {row['description']}{flag}"


def manual_text(row: dict) -> str:
    return f"{row['figure']} callout {row['callout']}: {row['key']}"


def part_cell(cell, definition: dict, titles: dict[str, str]) -> None:
    student = definition["student"]
    title = student["title"]["text"] + (f" ({student['context']})" if student.get("context") else "")
    line(cell, None, title, size=11, color=TEAL, first=True).runs[0].bold = True
    line(cell, None, f"{definition['partCount']} in the model · {definition['id']}", size=8, color=GREY)
    if student["aliases"]:
        line(cell, "Also called", ", ".join(student["aliases"]), size=8, color=GREY)
    old = definition["engineering"]["priorData"]["labels"]
    if old:
        line(cell, "Was", old[0], size=8, color=GREY)


def learner_cell(cell, definition: dict, titles: dict[str, str]) -> None:
    student = definition["student"]
    line(cell, "Does", student["function"]["text"], first=True)
    line(cell, "Made of", student["madeOf"]["text"])
    line(cell, "Sits", student["sits"]["text"])
    line(cell, "Works with", ", ".join(titles[i] for i in student["worksWith"]["ids"]))
    line(cell, "Look for", student["lookFor"]["text"])
    if student.get("catalogue"):
        line(cell, "Catalogue line", f"{student['catalogue']['description']}. {student['catalogue']['note']}", size=8, color=GREY)


def evidence_cell(cell, definition: dict) -> None:
    student = definition["student"]
    engineering = definition["engineering"]
    line(cell, "Name", evidence_text(student["title"]), size=8, first=True)
    for key, label in [("function", "Does"), ("madeOf", "Made of"), ("sits", "Sits"), ("worksWith", "Works with")]:
        line(cell, label, evidence_text(student[key]), size=8)
    alignment = engineering["alignment"]
    mark = RGBColor(0xB4, 0x23, 0x18) if alignment["result"] in ("mismatch", "not_found") else GREY
    line(cell, "Cross-check", f"{ALIGNMENT_WORDS[alignment['result']]}. {alignment['note']}", size=8, color=mark)
    for row in engineering["ipcIo520"]:
        line(cell, "IO-520 catalogue", ipc_text(row), size=8, color=GREY)
    for row in engineering["gtsioManual"]:
        line(cell, "GTSIO manual", manual_text(row), size=8, color=GREY)


def build(out_dir: Path) -> Path:
    doc = json.loads(DEFINITIONS.read_text(encoding="utf-8"))
    titles = {d["id"]: d["student"]["title"]["text"] + (f" ({d['student']['context']})" if d["student"].get("context") else "") for d in doc["definitions"]}
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))

    document.add_heading("GTSIO-520-H part details: draft text for your review", level=1)
    intro = document.add_paragraph()
    run(intro, "Tapping a part in the GTSIO-520-H whole-engine card will show a plain title, one sentence on what the part does, what it is made of, where it sits, "
               "what it works with, something to look for, and the catalogue wording. Every line below is a draft and stays marked 'Draft text' for students until you review it. "
               "Write in the last column; nothing here has been shown to students as reviewed.", size=10)
    rules = document.add_paragraph()
    run(rules, "How the text was written. ", bold=True, size=10)
    run(rules, "The title is the part noun from the IO-520 parts catalogue, in plain order and four words at most; where the catalogue has no row the GTSIO manual's figure key gives the name. "
               "Sizes and specifications are not in the plain lines: the 'catalogue line' carries the catalogue wording and is tagged as IO-520 (a similar engine) unless the GTSIO manual confirms it. "
               "No part numbers are shown to students (the GTSIO-520-H catalogue, Form X-30046A, is not available). Numbers appear only where a lesson or the manual gives them. "
               "Each line carries its evidence: 'stated in a source' and 'worked out from a source' name the source; 'reasoned, not in a source' gives the basis and is a claim to check; "
               "'look for' lines are prompts for the learner and claim nothing about the engine.", size=10)

    statuses = Counter()
    for definition in doc["definitions"]:
        for key in ("function", "madeOf", "sits"):
            statuses[(key, definition["student"][key]["status"])] += 1
    alignment = Counter(d["engineering"]["alignment"]["result"] for d in doc["definitions"])
    summary = document.add_paragraph()
    run(summary, f"{len(doc['definitions'])} kinds of part cover all 1,396 modelled parts. ", bold=True, size=10)
    run(summary, "Function lines: " + ", ".join(f"{statuses[('function', s)]} {STATUS_WORDS[s]}" for s in ("specified", "derived", "inferred")) + ". "
               "Catalogue cross-check: " + ", ".join(f"{alignment[a]} {a.replace('_', ' ')}" for a in ("confirmed", "gtsio_only", "io520_only", "mismatch", "not_found") if alignment[a]) + ".", size=10)

    flagged = [d for d in doc["definitions"] if d["engineering"]["alignment"]["result"] in ("mismatch", "not_found", "io520_only")]
    if flagged:
        document.add_heading("Where the two books disagree or one is silent", level=2)
        note = document.add_paragraph()
        run(note, "These are the entries to look at first. Each is in the table below with the full note.", size=10)
        for definition in flagged:
            paragraph = document.add_paragraph(style="List Bullet")
            run(paragraph, f"{titles[definition['id']]} ", bold=True, size=9)
            run(paragraph, f"[{definition['engineering']['alignment']['result'].replace('_', ' ')}] {definition['engineering']['alignment']['note']}", size=9)

    by_system: dict[str, list[dict]] = {}
    for definition in doc["definitions"]:
        by_system.setdefault(definition["student"]["system"], []).append(definition)
    for system in doc["systems"]:
        members = by_system.get(system["id"], [])
        if not members:
            continue
        document.add_page_break()
        document.add_heading(f"{system['label']} ({len(members)})", level=2)
        question = document.add_paragraph()
        run(question, system["question"], italic=True, size=10, color=GREY)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, head in zip(table.rows[0].cells, ["Part", "What a learner would read", "Where it comes from", "Your comments"]):
            shade(cell, HEADER_FILL)
            run(cell.paragraphs[0], head, bold=True, size=9)
        for definition in sorted(members, key=lambda d: titles[d["id"]].lower()):
            cells = table.add_row().cells
            part_cell(cells[0], definition, titles)
            learner_cell(cells[1], definition, titles)
            evidence_cell(cells[2], definition)
        set_widths(table, [3.6, 9.2, 9.6, 4.3])

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / FILE_NAME
    document.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--out", default=str(ROOT.parent / "Lesson Plans" / "Review sheets"), help="folder for the sheet")
    args = parser.parse_args()
    print(build(Path(args.out)))


if __name__ == "__main__":
    main()
