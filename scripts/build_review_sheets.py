"""Write one Word review sheet per numbered lesson, for the instructor to read through and mark up.

A sheet lists everything a student meets in the lesson: each step's text, what it shows (picture, 3D pose or interactive), the source
note behind it, and every check with its answers and rationale, each beside an empty comments column. The sheets are generated from the
lesson packs in web/ and are not committed; run this again after the packs change.

    python scripts/build_review_sheets.py [--out "../Lesson Plans/Review sheets"] [--only lesson-id ...]
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
WEB = ROOT / "web"
TEAL = RGBColor(0x0F, 0x76, 0x6E)
GREY = RGBColor(0x55, 0x55, 0x55)
HEADER_FILL = "D9EFEC"


def load_packs() -> list[dict]:
    manifest = json.loads((WEB / "lessons-manifest.json").read_text(encoding="utf-8"))
    packs = [json.loads((WEB / path.removeprefix("./")).read_text(encoding="utf-8")) for path in manifest["packs"]]
    for pack, path in zip(packs, manifest["packs"]):
        pack["_file"] = path.removeprefix("./")
    return packs


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_widths(table, widths_cm: list[float]) -> None:
    """Fix the column widths in the table grid as well as in each cell: Word follows the cells, LibreOffice follows the grid."""
    table.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def write(cell, text: str, *, bold: bool = False, colour: RGBColor | None = None, size: float = 9, italic: bool = False) -> None:
    paragraph = cell.paragraphs[0] if not cell.paragraphs[0].text and len(cell.paragraphs) == 1 else cell.add_paragraph()
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)
    if colour:
        run.font.color.rgb = colour
    paragraph.paragraph_format.space_after = Pt(2)


def add_picture(cell, relative_url: str, width_cm: float = 4.9) -> bool:
    path = WEB / relative_url.removeprefix("./")
    if not path.exists() or path.suffix.lower() not in {".webp", ".png", ".jpg", ".jpeg"}:
        return False
    image = Image.open(path).convert("RGB")
    image.thumbnail((900, 900))
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    buffer.seek(0)
    cell.add_paragraph().add_run().add_picture(buffer, width=Cm(width_cm))
    return True


def media_lines(step: dict) -> list[str]:
    """What the student sees above the step's text, in words."""
    plan = step.get("mediaPlan") or {}
    url = step.get("url") or ""
    if plan.get("mode") == "none":
        return ["Text only"]
    if url.startswith("artifact:"):
        return [f"Interactive: {url.removeprefix('artifact:')}"]
    if step.get("type") == "model-pose":
        lines = [f"3D model: {step.get('modelId')}"]
        if step.get("savedMotionId"):
            lines.append(f"Saved motion: {step['savedMotionId']}" + (f" at {step['motionProgress']}%" if step.get("motionProgress") is not None else ""))
        if step.get("action"):
            lines.append(f"Pose: {step['action'].get('type')} {step['action'].get('value')}")
        if step.get("viewPreset"):
            lines.append(f"View: {step['viewPreset']}")
        if step.get("focusParts"):
            lines.append(f"Spotlight ({step.get('focusMode', 'x-ray')}): {', '.join(step['focusParts'])}")
        return lines
    if step.get("type") == "image" and url:
        lines = [f"Picture: {step.get('alt') or url}"]
        if step.get("credit"):
            lines.append(f"Credit: {step['credit']}" + (f" ({step['license']})" if step.get("license") else ""))
        return lines
    if step.get("type") == "web-embed":
        return [f"Embedded page: {url}"]
    return ["Text only"]


def add_steps(document: Document, lesson: dict) -> None:
    document.add_heading("Steps", level=2)
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ["#", "Step", "What the student reads", "What it shows", "Source note", "Your comments"]):
        write(cell, label, bold=True, size=9)
        shade(cell, HEADER_FILL)
    repeat_header(table.rows[0])
    for number, step in enumerate(lesson["steps"], start=1):
        row = table.add_row()
        cells = row.cells
        write(cells[0], str(number), bold=True)
        write(cells[1], step.get("title") or f"Step {number}", bold=True)
        write(cells[1], step.get("type", "text"), colour=GREY, size=8)
        if step.get("deepDiveLinks"):
            write(cells[1], "Deep dive: " + ", ".join(step["deepDiveLinks"]), colour=GREY, size=8)
        write(cells[2], step.get("prompt", ""))
        for line in media_lines(step):
            write(cells[3], line, size=8.5)
        if step.get("url") and step.get("type") == "image":
            add_picture(cells[3], step["url"])
        if step.get("note"):
            write(cells[4], step["note"], size=8.5, colour=GREY)
        if step.get("sourceRefs"):
            write(cells[4], "Sources: " + ", ".join(step["sourceRefs"]), size=8, colour=GREY, italic=True)
    set_widths(table, [0.8, 3.0, 8.0, 5.4, 5.0, 4.4])


def add_checks(document: Document, checks: list[dict]) -> None:
    document.add_heading("Checks", level=2)
    if not checks:
        document.add_paragraph("This lesson has no check questions yet.")
        return
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, label in zip(table.rows[0].cells, ["#", "Question", "Answers (correct one marked)", "Why, and the hint", "Your comments"]):
        write(cell, label, bold=True, size=9)
        shade(cell, HEADER_FILL)
    repeat_header(table.rows[0])
    for number, check in enumerate(checks, start=1):
        cells = table.add_row().cells
        write(cells[0], str(number), bold=True)
        write(cells[1], check["question"])
        write(cells[1], check["type"], colour=GREY, size=8)
        if check["type"] == "ordering":
            for position, item in enumerate(check.get("items", []), start=1):
                write(cells[2], f"{position}. {item}", size=8.5)
        else:
            for index, answer in enumerate(check.get("answers", [])):
                correct = index == check.get("correct")
                write(cells[2], f"{'✔ ' if correct else '   '}{chr(65 + index)}. {answer}", bold=correct, size=8.5)
        write(cells[3], check.get("rationale", ""), size=8.5)
        if check.get("hint"):
            write(cells[3], "Hint: " + check["hint"], size=8.5, colour=GREY, italic=True)
    set_widths(table, [0.8, 6.4, 8.0, 7.0, 4.4])


def add_sources(document: Document, pack: dict, lesson: dict) -> None:
    cited = {ref for step in lesson["steps"] for ref in step.get("sourceRefs", [])}
    sources = [source for source in pack.get("sources", []) if source["id"] in cited]
    if not sources:
        return
    document.add_heading("Sources this lesson cites", level=2)
    for source in sources:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(f"{source['title']}")
        run.bold = True
        run.font.size = Pt(9)
        detail = paragraph.add_run(f"  (tier {source['tier']}) {source.get('applicability', '')}")
        detail.font.size = Pt(8.5)
        detail.font.color.rgb = GREY


def build_sheet(pack: dict, lesson: dict, out: Path) -> Path:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.5)
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(10)

    number = lesson.get("sequenceNumber")
    heading = document.add_heading(f"Lesson {number:02d} · {lesson['title']}" if number else lesson["title"], level=1)
    heading.runs[0].font.color.rgb = TEAL

    meta = document.add_table(rows=0, cols=2)
    meta.style = "Table Grid"
    status = lesson.get("reviewStatus", "unreviewed")
    rows = [
        ("Course", pack["title"]),
        ("Review status", status + (f" on {lesson['reviewedOn']}" if lesson.get("reviewedOn") else "")),
        ("3D models", ", ".join(lesson.get("models") or []) or "none"),
        ("Steps and checks", f"{len(lesson['steps'])} steps, {len([c for c in pack.get('checks', []) if c['lessonId'] == lesson['id']])} checks"),
    ]
    for label, value in rows:
        cells = meta.add_row().cells
        write(cells[0], label, bold=True)
        shade(cells[0], HEADER_FILL)
        write(cells[1], value)
    set_widths(meta, [4.0, 22.0])

    document.add_heading("Objective", level=2)
    document.add_paragraph(lesson.get("objective", ""))
    if lesson.get("completionCriteria"):
        document.add_heading("Completion criteria", level=2)
        document.add_paragraph(lesson["completionCriteria"])

    add_sources(document, pack, lesson)
    add_steps(document, lesson)
    add_checks(document, [check for check in pack.get("checks", []) if check["lessonId"] == lesson["id"]])

    document.add_heading("Your decision", level=2)
    for line in ["☐  Good as it is", "☐  Good after the edits in the comments column", "☐  Needs reworking", "Notes:"]:
        document.add_paragraph(line)
    note = document.add_paragraph()
    run = note.add_run(
        "To record a decision, tell Claude which lesson and the date, or set reviewStatus to \"reviewed\" and reviewedOn to YYYY-MM-DD "
        f"for this lesson in web/{pack['_file']}. Until then students see an \"Instructor review pending\" chip."
    )
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY

    name = re.sub(r'[<>:"/\\|?*]', "", f"Lesson {number:02d} - {lesson['title']}" if number else lesson["title"])
    target = out / f"{name}.docx"
    document.save(target)
    return target


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=ROOT.parent / "Lesson Plans" / "Review sheets", help="folder for the .docx files")
    parser.add_argument("--only", nargs="*", help="lesson ids to build (default: every numbered lesson)")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    built = 0
    for pack in load_packs():
        for lesson in pack["lessons"]:
            if args.only and lesson["id"] not in args.only:
                continue
            if not args.only and not lesson.get("sequenceNumber"):
                continue
            print(build_sheet(pack, lesson, args.out).name)
            built += 1
    print(f"{built} review sheets in {args.out}")
    return 0 if built else 1


if __name__ == "__main__":
    sys.exit(main())
